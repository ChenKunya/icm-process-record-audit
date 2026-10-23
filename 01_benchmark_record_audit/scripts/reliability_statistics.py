"""Inter-coder reliability on the audit subsample (8 benchmarks x 10 fields = 80 units).

Usage:
    python scripts/reliability_statistics.py [coding_table.xlsx] [form_A02.xlsx] [form_A03.xlsx] [output.json]

Defaults: the files in data/ and results/reliability_statistics.json.
Requires Python 3.9+, numpy and openpyxl.

A01 is the primary coder (full coding table, Protocol A codes). A02 and A03 are the two
annotators. Their codes are derived from the three answers on the "Annotation" sheet:
    ① No                                         -> NA
    ① Yes, ② Directly visible                    -> D
    ① Yes, ② Not visible, ③ Uniquely recoverable -> R
    ① Yes, ② Not visible, ③ Not recoverable      -> U

Reported: Krippendorff's alpha (nominal) for the three coders and for A02 and A03, with
95% percentile intervals from 10,000 bootstrap resamples of benchmarks (seed 20261020);
agreement and Cohen's kappa for each pair; alpha for the applicability stage and for
record status among units all three coders judged applicable; Class I/C/M proxies on
the subsample for each coder; and the list of disagreeing units.
"""
import collections
import itertools
import json
import sys
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / 'data'
ARGS = sys.argv[1:] + [None] * 4
TABLE = Path(ARGS[0] or DATA / 'audit_coding_table_74x10_A01.xlsx')
FORMS = {'A02': Path(ARGS[1] or DATA / 'reliability_annotation_A02.xlsx'),
         'A03': Path(ARGS[2] or DATA / 'reliability_annotation_A03.xlsx')}
DST = Path(ARGS[3] or ROOT / 'results' / 'reliability_statistics.json')

SEED, N_BOOT = 20261020, 10000
FIELDS = ['Source identifier', 'Temporal position', 'State version', 'Dependence relations',
          'Parameter provenance', 'Propagable uncertainty representation', 'Conflict representation',
          'Responsibility or commitment status', 'Validity-change event', 'Version lineage and rollback']
CODERS = ['A01', 'A02', 'A03']


def text(x):
    return '' if x is None else str(x).strip()


def read_primary(path):
    ws = openpyxl.load_workbook(path)['Coding']
    codes, ids = {}, {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[1] is None:
            continue
        b, name, field = int(r[1]), text(r[2]), text(r[4])
        ids[name] = b
        codes[(b, FIELDS.index(field))] = 'NA' if text(r[5]) == 'Not applicable' else text(r[7])
    return codes, ids


def read_form(path, ids):
    ws = openpyxl.load_workbook(path)['Annotation']
    codes, incomplete = {}, []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        key = (ids[text(r[1])], FIELDS.index(text(r[2])))
        q1, q2, q3 = text(r[3]), text(r[4]), text(r[5])
        code = None
        if q1 == 'No':
            code = 'NA'
        elif q1 == 'Yes' and q2 == 'Directly visible':
            code = 'D'
        elif q1 == 'Yes' and q2 == 'Not visible':
            code = {'Uniquely recoverable': 'R', 'Not recoverable': 'U'}.get(q3)
        if code is None:
            incomplete.append(key)
        codes[key] = code
    return codes, incomplete


def pct(k, n, digits=1):
    """Percentage rounded half up, as reported in the manuscript."""
    q = Decimal(1).scaleb(-digits)
    return float((Decimal(100) * k / n).quantize(q, rounding=ROUND_HALF_UP))


def kalpha(units):
    """Krippendorff's alpha with the nominal metric. Each unit is a list of values (None = missing)."""
    o = collections.Counter()
    for vals in units:
        v = [x for x in vals if x is not None]
        if len(v) < 2:
            continue
        for a, b in itertools.permutations(range(len(v)), 2):
            o[(v[a], v[b])] += 1.0 / (len(v) - 1)
    cats = sorted({c for c, _ in o})
    nc = {c: sum(o[(c, k)] for k in cats) for c in cats}
    n = sum(nc.values())
    d_obs = sum(o[(c, k)] for c in cats for k in cats if c != k)
    d_exp = sum(nc[c] * nc[k] for c in cats for k in cats if c != k)
    if d_exp == 0:
        return 1.0 if d_obs == 0 else float('nan')
    return 1.0 - (n - 1) * d_obs / d_exp


def cohen(a, b):
    n = len(a)
    po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = collections.Counter(a), collections.Counter(b)
    pe = sum(ca[c] * cb[c] for c in set(ca) | set(cb)) / n / n
    return dict(n=n, agreement_pct=pct(sum(x == y for x, y in zip(a, b)), n), kappa=round((po - pe) / (1 - pe), 2))


def proxies(codes, bench):
    ok = lambda v: v in ('D', 'R')
    i_k = sum(ok(codes[(b, 0)]) and ok(codes[(b, 1)]) for b in bench)
    c_base = [b for b in bench if codes[(b, 3)] != 'NA' and codes[(b, 4)] != 'NA']
    m_base = [b for b in bench if codes[(b, 8)] != 'NA']
    return dict(class_I=f'{i_k}/{len(bench)}',
                class_C=f"{sum(ok(codes[(b, 3)]) and ok(codes[(b, 4)]) for b in c_base)}/{len(c_base)}",
                class_M=f'{sum(ok(codes[(b, 8)]) for b in m_base)}/{len(m_base)}')


def main():
    primary, ids = read_primary(TABLE)
    coder = {}
    incomplete = {}
    for cid, path in FORMS.items():
        coder[cid], incomplete[cid] = read_form(path, ids)
    units = sorted(coder['A02'])
    assert sorted(coder['A03']) == units and len(units) == 80
    coder['A01'] = {u: primary[u] for u in units}
    bench = sorted({b for b, _ in units})
    names = {b: n for n, b in ids.items()}

    def alpha(us, who, f=lambda x: x, keep=lambda u: True):
        return kalpha([[f(coder[c][u]) for c in who] for u in us if keep(u)])

    pair = ['A02', 'A03']
    out = dict(n_units=len(units), benchmarks=[names[b] for b in bench],
               incomplete_units={k: len(v) for k, v in incomplete.items()})
    out['alpha_A01_A02_A03'] = round(alpha(units, CODERS), 2)
    out['alpha_A02_A03'] = round(alpha(units, pair), 2)

    rng = np.random.default_rng(SEED)
    by_b = {b: [u for u in units if u[0] == b] for b in bench}
    boot3, boot2 = [], []
    for _ in range(N_BOOT):
        us = [u for i in rng.integers(0, len(bench), len(bench)) for u in by_b[bench[i]]]
        boot3.append(alpha(us, CODERS))
        boot2.append(alpha(us, pair))
    ci = lambda x: [round(float(np.nanpercentile(x, q)), 2) for q in (2.5, 97.5)]
    out['alpha_A01_A02_A03_bootstrap95'] = ci(boot3)
    out['alpha_A02_A03_bootstrap95'] = ci(boot2)

    out['pairwise'] = {f'{a}_{b}': cohen([coder[a][u] for u in units], [coder[b][u] for u in units])
                       for a, b in itertools.combinations(CODERS, 2)}

    applicable = lambda x: 'NA' if x == 'NA' else 'applicable'
    all_app = lambda u: all(coder[c][u] != 'NA' for c in CODERS)
    out['alpha_applicability_stage'] = round(alpha(units, CODERS, applicable), 2)
    out['n_units_applicable_for_all'] = sum(all_app(u) for u in units)
    out['alpha_status_stage'] = round(alpha(units, CODERS, keep=all_app), 2)
    out['status_agreement_with_A01'] = {
        c: dict(n_both_applicable=sum(coder['A01'][u] != 'NA' and coder[c][u] != 'NA' for u in units),
                n_same_status=sum(coder['A01'][u] != 'NA' and coder[c][u] == coder['A01'][u] for u in units))
        for c in pair}
    out['field_all_three_agree'] = {FIELDS[j]: sum(len({coder[c][(b, j)] for c in CODERS}) == 1 for b in bench)
                                    for j in range(len(FIELDS))}
    out['proxies_on_subsample'] = {c: proxies(coder[c], bench) for c in CODERS}
    out['confusion_rows_A01'] = {c: {a: {b: sum(coder['A01'][u] == a and coder[c][u] == b for u in units)
                                         for b in ('D', 'R', 'U', 'NA')} for a in ('D', 'R', 'U', 'NA')}
                                 for c in pair}
    out['disagreements'] = [dict(benchmark=names[b], field=FIELDS[j], **{c: coder[c][(b, j)] for c in CODERS})
                            for b, j in units if len({coder[c][(b, j)] for c in CODERS}) > 1]

    DST.parent.mkdir(parents=True, exist_ok=True)
    DST.write_text(json.dumps(out, indent=1), encoding='utf-8')
    print(f"alpha (A01, A02, A03) = {out['alpha_A01_A02_A03']} {out['alpha_A01_A02_A03_bootstrap95']}")
    print(f'Written to {DST}')


if __name__ == '__main__':
    main()
